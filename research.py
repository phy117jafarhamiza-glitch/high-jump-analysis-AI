"""
أدوات البحث العلمي:
  1) الصدق والثبات (Pearson، التحيز، حدود الاتفاق، ICC، SEM، MDC)
  2) جداول الفصل الرابع (التجانس، الاعتدالية، التكافؤ، الفروق، نسب التطور)
  3) نسب اتفاق الخبراء وكا² وتحديث ملف الأنموذج
  4) تصدير كل ذلك إلى Word بجداول من اليمين لليسار
"""
import io
import math

import numpy as np
import pandas as pd
from scipy import stats

import model

ALPHA = 0.05


def _f(v, nd=2):
    try:
        x = float(v)
    except (TypeError, ValueError):
        return str(v) if v is not None else "—"
    if math.isnan(x):
        return "—"
    return f"{x:.{nd}f}"


def _sig(p):
    if p is None or (isinstance(p, float) and math.isnan(p)):
        return "—"
    return "دال" if p < ALPHA else "غير دال"


# ============================================================================
# Word
# ============================================================================
class WordReport:
    def __init__(self, title=None, landscape=False):
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt

        self.doc = Document()
        self.n = 0
        sec = self.doc.sections[0]
        if landscape:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = sec.page_height, sec.page_width
        sec.left_margin = sec.right_margin = Cm(2)
        b = OxmlElement("w:bidi"); b.set(qn("w:val"), "1"); sec._sectPr.append(b)
        st = self.doc.styles["Normal"]
        st.font.name = "Simplified Arabic"
        st.font.size = Pt(12)
        if title:
            model._para(self.doc, title, bold=True, size=15, align="center")

    def text(self, t, bold=False, size=12):
        model._para(self.doc, t, bold=bold, size=size)

    def table(self, caption, df, note=None, start_no=None):
        from docx.oxml import OxmlElement
        from docx.shared import Pt

        self.n = start_no if start_no else self.n + 1
        model._para(self.doc, f"جدول ({self.n}) {caption}", bold=True, size=12, align="center")
        t = self.doc.add_table(rows=1, cols=len(df.columns))
        t.style = "Table Grid"
        model._table_rtl(t)
        for i, c in enumerate(df.columns):
            p = t.rows[0].cells[i].paragraphs[0]
            r = p.add_run(str(c)); r.bold = True; r.font.size = Pt(10)
            r._r.get_or_add_rPr().append(OxmlElement("w:bCs"))
            p.alignment = 1
            model._rtl_paragraph(p)
        for _, row in df.iterrows():
            cells = t.add_row().cells
            for i, v in enumerate(row):
                p = cells[i].paragraphs[0]
                txt = "—" if v is None or (isinstance(v, float) and math.isnan(v)) else str(v)
                if any(ch.isdigit() for ch in txt) and "-" in txt:
                    # علامة اتجاه من اليسار لليمين حتى تظهر الإشارة السالبة قبل الرقم
                    txt = " ".join("\u200e" + w if any(c.isdigit() for c in w) else w for w in txt.split(" "))
                r = p.add_run(txt)
                r.font.size = Pt(10)
                p.alignment = 1
                model._rtl_paragraph(p)
        # عرض الأعمدة: الأول (أسماء المتغيرات/البنود) أعرض
        from docx.shared import Cm
        sec = self.doc.sections[0]
        usable = (sec.page_width - sec.left_margin - sec.right_margin) / 360000  # سم
        ncol = len(df.columns)
        first = min(7.0, usable * 0.28) if ncol > 3 else usable / ncol
        rest = (usable - first) / max(1, ncol - 1)
        t.autofit = False
        widths = [first] + [rest] * (ncol - 1)
        if df.columns[0] == "ت" and ncol > 2:
            rest = (usable - first - 1.0) / (ncol - 2)
            widths = [1.0, first] + [rest] * (ncol - 2)
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(widths[i] * 567)))
        for row in t.rows:
            for i, c in enumerate(row.cells):
                c.width = Cm(widths[i])
        if note:
            model._para(self.doc, note, size=9)
        model._para(self.doc, "")

    def bytes(self):
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()


# ============================================================================
# 1) الصدق والثبات
# ============================================================================
VR_COLS = ["رمز_المحاولة", "المتغير", "القياس_أ", "القياس_ب"]
VR_MODES = {
    "validity": ("الصدق (أ = التطبيق، ب = Kinovea أو المحك)", "التطبيق", "المحك"),
    "retest_analysis": ("ثبات إعادة التحليل (أ = التحليل الأول، ب = الثاني)", "التحليل الأول", "التحليل الثاني"),
    "retest_test": ("ثبات إعادة الاختبار (أ = الاختبار الأول، ب = الثاني)", "الاختبار الأول", "الاختبار الثاني"),
}


def icc(a, b):
    """ICC(2,1) اتفاق مطلق و ICC(3,1) اتساق لقياسين (McGraw & Wong 1996)."""
    Y = np.c_[a, b].astype(float)
    n, k = Y.shape
    gm = Y.mean()
    ssr = k * ((Y.mean(1) - gm) ** 2).sum()
    ssc = n * ((Y.mean(0) - gm) ** 2).sum()
    sst = ((Y - gm) ** 2).sum()
    sse = sst - ssr - ssc
    msr, msc, mse = ssr / (n - 1), ssc / (k - 1), sse / ((n - 1) * (k - 1))
    icc21 = (msr - mse) / (msr + (k - 1) * mse + k * (msc - mse) / n) if (msr + (k - 1) * mse + k * (msc - mse) / n) else np.nan
    icc31 = (msr - mse) / (msr + (k - 1) * mse) if (msr + (k - 1) * mse) else np.nan
    # فترة الثقة 95% لـ ICC(3,1)
    lo = hi = np.nan
    if mse > 0 and n > 2:
        F = msr / mse
        fl = F / stats.f.ppf(0.975, n - 1, (n - 1) * (k - 1))
        fu = F * stats.f.ppf(0.975, (n - 1) * (k - 1), n - 1)
        lo, hi = (fl - 1) / (fl + k - 1), (fu - 1) / (fu + k - 1)
    return icc21, icc31, lo, hi


def icc_label(v):
    if v is None or np.isnan(v):
        return "—"
    return "ضعيف" if v < 0.5 else "متوسط" if v < 0.75 else "جيد" if v < 0.9 else "ممتاز"


def vr_template(sample=False):
    rows = []
    if sample:
        rng = np.random.default_rng(3)
        for var, m, sd, err in (("knee_at_plant_deg", 160, 6, 2.0), ("contact_time_s", 0.19, 0.02, 0.008),
                                ("takeoff_angle_deg", 45, 4, 1.5), ("trunk_lean_takeoff_deg", 4, 5, 1.8)):
            for i in range(1, 13):
                truth = m + sd * rng.normal()
                rows.append([f"T{i:02d}", var, round(truth + err * rng.normal(), 3), round(truth + err * 0.6 * rng.normal(), 3)])
    df = pd.DataFrame(rows, columns=VR_COLS)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="البيانات", index=False)
        pd.DataFrame({"رمز المتغير": list(model.METRIC_LABELS),
                      "الاسم": [v[0] for v in model.METRIC_LABELS.values()],
                      "الوحدة": [v[1] for v in model.METRIC_LABELS.values()]}).to_excel(w, sheet_name="رموز المتغيرات", index=False)
        pd.DataFrame({"تعليمات": [
            "صف لكل محاولة ولكل متغير.",
            "الصدق: القياس_أ = قيمة التطبيق، القياس_ب = قيمة Kinovea (أو المحك) للمحاولة نفسها.",
            "ثبات إعادة التحليل: القياس_أ = التحليل الأول، القياس_ب = إعادة تحليل الفيديو نفسه.",
            "ثبات إعادة الاختبار: القياس_أ = الاختبار الأول، القياس_ب = الاختبار الثاني بعد أيام.",
            "يُفضَّل 10 محاولات فأكثر لكل متغير.",
            "عمود المتغير: رمز المتغير (مثل knee_at_plant_deg) أو اسمه كما في ورقة «رموز المتغيرات».",
        ] + (["⚠️ بيانات تجريبية مولّدة لتجربة الصفحة فقط"] if sample else [])}).to_excel(w, sheet_name="اقرأني", index=False)
    return buf.getvalue()


def vr_read(file_bytes):
    df = pd.read_excel(io.BytesIO(file_bytes), sheet_name=0)
    missing = [c for c in VR_COLS[1:] if c not in df.columns]
    if missing:
        raise ValueError("أعمدة ناقصة: " + "، ".join(missing))
    name_to_code = {v[0]: k for k, v in model.METRIC_LABELS.items()}
    df["المتغير"] = df["المتغير"].astype(str).str.strip().map(lambda v: name_to_code.get(v, v))
    for c in ("القياس_أ", "القياس_ب"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    return df.dropna(subset=["القياس_أ", "القياس_ب"])


def vr_analyze(df, mode):
    rows = []
    for var, g in df.groupby("المتغير", sort=False):
        a, b = g["القياس_أ"].to_numpy(float), g["القياس_ب"].to_numpy(float)
        n = len(a)
        lab, unit = model.METRIC_LABELS.get(var, (var, ""))
        r = p = np.nan
        if n >= 3 and np.std(a) > 0 and np.std(b) > 0:
            r, p = stats.pearsonr(a, b)
        d = a - b
        bias, sdd = d.mean(), d.std(ddof=1) if n > 1 else np.nan
        i21, i31, lo, hi = icc(a, b) if n >= 3 else (np.nan,) * 4
        sd_all = np.std(np.r_[a, b], ddof=1) if n > 1 else np.nan
        rel = i21 if mode == "validity" else i31 if mode == "retest_analysis" else i21
        sem = sd_all * math.sqrt(max(0.0, 1 - rel)) if not np.isnan(rel) else np.nan
        rows.append({
            "المتغير": lab, "الوحدة": unit, "ن": n,
            f"الوسط ({VR_MODES[mode][1]})": a.mean(), f"الوسط ({VR_MODES[mode][2]})": b.mean(),
            "معامل الارتباط (ر)": r, "مستوى الدلالة": p,
            "متوسط الفرق (التحيز)": bias, "متوسط الفرق المطلق": np.abs(d).mean(),
            "حدود الاتفاق 95%": (f"{_f(bias - 1.96 * sdd)} إلى {_f(bias + 1.96 * sdd)}" if not np.isnan(sdd) else "—"),
            "ICC(2,1)": i21, "ICC(3,1)": i31, "فترة الثقة 95% لـ ICC(3,1)": (f"{_f(lo)} – {_f(hi)}" if not np.isnan(lo) else "—"),
            "SEM": sem, "MDC95": sem * 1.96 * math.sqrt(2) if not np.isnan(sem) else np.nan,
            "التقدير": icc_label(rel),
        })
    return pd.DataFrame(rows)


def vr_word(res, mode, table_no=None):
    title, la, lb = VR_MODES[mode]
    rep = WordReport(landscape=True)
    if mode == "validity":
        cols = ["المتغير", "الوحدة", "ن", f"الوسط ({la})", f"الوسط ({lb})", "معامل الارتباط (ر)", "مستوى الدلالة",
                "متوسط الفرق المطلق", "حدود الاتفاق 95%", "ICC(2,1)"]
        cap = "يبين معاملات صدق المحك للمتغيرات البايوميكانيكية"
        note = f"المحك: {lb}. حدود الاتفاق وفق طريقة Bland–Altman. مستوى الدلالة (0.05)."
    else:
        key = "ICC(3,1)" if mode == "retest_analysis" else "ICC(2,1)"
        cols = ["المتغير", "الوحدة", "ن", f"الوسط ({la})", f"الوسط ({lb})", key, "فترة الثقة 95% لـ ICC(3,1)", "SEM", "MDC95", "التقدير"]
        cap = "يبين معاملات ثبات " + ("إعادة التحليل" if mode == "retest_analysis" else "إعادة الاختبار")
        note = "تفسير ICC: أقل من 0.50 ضعيف، 0.50–0.75 متوسط، 0.75–0.90 جيد، أكبر من 0.90 ممتاز (Koo & Li, 2016). SEM: الخطأ المعياري للقياس، MDC95: أقل تغير يمكن اكتشافه."
    d = res[cols].copy()
    for c in d.columns:
        if d[c].dtype.kind == "f":
            d[c] = d[c].map(lambda v: _f(v, 3 if "ICC" in c or "الدلالة" in c or c == "معامل الارتباط (ر)" else 2))
    rep.table(cap, d, note, start_no=table_no)
    return rep.bytes()


# ============================================================================
# 2) جداول الفصل الرابع
# ============================================================================
def _skew_pearson(x):
    x = np.asarray(x, float)
    x = x[~np.isnan(x)]
    if len(x) < 3 or x.std(ddof=1) == 0:
        return np.nan
    return 3 * (x.mean() - np.median(x)) / x.std(ddof=1)


def homogeneity(players, long_df, pre_achievement_var="A01"):
    rows = []
    p = players.copy()
    this_year = pd.Timestamp.now().year
    specs = [("الطول", "سم", pd.to_numeric(p.get("height"), errors="coerce")),
             ("الكتلة", "كغم", pd.to_numeric(p.get("weight"), errors="coerce"))]
    if "birth_year" in p:
        specs.append(("العمر الزمني", "سنة", this_year - pd.to_numeric(p["birth_year"], errors="coerce")))
    if "training_years" in p:
        specs.append(("العمر التدريبي", "سنة", pd.to_numeric(p["training_years"], errors="coerce")))
    ach = long_df[(long_df["var"] == pre_achievement_var) & (long_df["phase"] == "قبلي")]["value"] if not long_df.empty else []
    if len(ach):
        specs.append(("الإنجاز (قبلي)", "م", ach))
    for name, unit, x in specs:
        x = pd.Series(x).dropna().astype(float)
        if x.empty:
            continue
        rows.append({"المتغيرات": name, "وحدة القياس": unit, "الوسط الحسابي": _f(x.mean()), "الوسيط": _f(x.median()),
                     "الانحراف المعياري": _f(x.std(ddof=1)), "معامل الالتواء": _f(_skew_pearson(x), 3)})
    return pd.DataFrame(rows)


def normality(long_df, phase="قبلي"):
    rows = []
    d = long_df[long_df["phase"] == phase]
    for (var, lab, unit), g in d.groupby(["var", "label", "unit"], sort=False):
        x = g["value"].dropna().to_numpy(float)
        w = p = np.nan
        if len(x) >= 3 and np.std(x) > 0:
            w, p = stats.shapiro(x)
        rows.append({"المتغيرات": lab, "الوحدة": unit, "ن": len(x), "قيمة (W)": _f(w, 3), "مستوى الدلالة": _f(p, 3),
                     "التوزيع": "—" if np.isnan(p) else ("طبيعي" if p >= ALPHA else "غير طبيعي")})
    return pd.DataFrame(rows)


def _ms(x):
    x = np.asarray(x, float)
    x = x[~np.isnan(x)]
    return (x.mean() if len(x) else np.nan), (x.std(ddof=1) if len(x) > 1 else np.nan), len(x)


def between_groups(long_df, phase, groups=("تجريبية", "ضابطة")):
    rows = []
    d = long_df[long_df["phase"] == phase]
    for (var, lab, unit), g in d.groupby(["var", "label", "unit"], sort=False):
        a = g[g["group"] == groups[0]]["value"].to_numpy(float)
        b = g[g["group"] == groups[1]]["value"].to_numpy(float)
        m1, s1, n1 = _ms(a)
        m2, s2, n2 = _ms(b)
        t = p = np.nan
        if n1 >= 2 and n2 >= 2 and (np.std(a) > 0 or np.std(b) > 0):
            t, p = stats.ttest_ind(a, b)
        rows.append({"المتغيرات": lab, "الوحدة": unit,
                     f"{groups[0]}: س": _f(m1), f"{groups[0]}: ع": _f(s1), f"{groups[1]}: س": _f(m2), f"{groups[1]}: ع": _f(s2),
                     "قيمة (t) المحسوبة": _f(t, 3), "مستوى الدلالة": _f(p, 3), "الدلالة": _sig(p)})
    return pd.DataFrame(rows)


def within_group(long_df, group):
    rows = []
    d = long_df[long_df["group"] == group]
    for (var, lab, unit, better), g in d.groupby(["var", "label", "unit", "better"], sort=False):
        piv = g.pivot_table(index="code", columns="phase", values="value", aggfunc="first")
        if not {"قبلي", "بعدي"} <= set(piv.columns):
            continue
        piv = piv.dropna(subset=["قبلي", "بعدي"])
        m1, s1, n = _ms(piv["قبلي"])
        m2, s2, _ = _ms(piv["بعدي"])
        diff = (piv["بعدي"] - piv["قبلي"]).to_numpy(float)
        fm, fs = (diff.mean(), diff.std(ddof=1)) if n > 1 else (np.nan, np.nan)
        t = p = np.nan
        if n >= 2 and fs and fs > 0:
            t, p = stats.ttest_rel(piv["بعدي"], piv["قبلي"])
        rows.append({"المتغيرات": lab, "الوحدة": unit, "قبلي: س": _f(m1), "قبلي: ع": _f(s1), "بعدي: س": _f(m2), "بعدي: ع": _f(s2),
                     "س ف": _f(fm), "ع ف": _f(fs), "قيمة (t) المحسوبة": _f(t, 3), "مستوى الدلالة": _f(p, 3), "الدلالة": _sig(p),
                     "حجم الأثر (d)": _f(fm / fs if fs else np.nan, 2)})
    return pd.DataFrame(rows)


def improvement(long_df, groups=("تجريبية", "ضابطة")):
    rows = []
    for (var, lab, unit, better), g in long_df.groupby(["var", "label", "unit", "better"], sort=False):
        row = {"المتغيرات": lab, "الوحدة": unit}
        for gr in groups:
            gg = g[g["group"] == gr]
            pre, post = gg[gg["phase"] == "قبلي"]["value"].mean(), gg[gg["phase"] == "بعدي"]["value"].mean()
            ok = pd.notna(pre) and pd.notna(post) and pre != 0 and not (pre < 0 or post < 0)
            v = (post - pre) / abs(pre) * 100 * (-1 if better == "أقل" else 1) if ok else np.nan
            row[f"نسبة التطور % ({gr})"] = _f(v)
        rows.append(row)
    return pd.DataFrame(rows)


def chapter4_word(players, long_df, groups=("تجريبية", "ضابطة"), start_no=1):
    rep = WordReport("جداول الفصل الرابع (مسودة آلية للمراجعة)", landscape=True)
    rep.n = start_no - 1
    note_sig = "مستوى الدلالة (0.05). س: الوسط الحسابي، ع: الانحراف المعياري."
    t = homogeneity(players, long_df)
    if not t.empty:
        rep.table("يبين تجانس عينة البحث في المتغيرات المختارة", t, "معامل الالتواء = 3 (الوسط − الوسيط) ÷ الانحراف المعياري؛ التجانس عندما يقع ضمن (±1).")
    t = normality(long_df, "قبلي")
    if not t.empty:
        rep.table("يبين نتائج اختبار (Shapiro–Wilk) لاعتدالية التوزيع في الاختبار القبلي", t, "التوزيع طبيعي عندما يكون مستوى الدلالة أكبر من (0.05).")
    t = between_groups(long_df, "قبلي", groups)
    if not t.empty:
        rep.table("يبين تكافؤ مجموعتي البحث في الاختبار القبلي", t, note_sig + " التكافؤ عندما يكون الفرق غير دال.")
    for gr in groups:
        t = within_group(long_df, gr)
        if not t.empty:
            rep.table(f"يبين الفروق بين الاختبارين القبلي والبعدي للمجموعة {gr}", t,
                      note_sig + " س ف: متوسط الفروق، ع ف: انحراف الفروق. اختبار (t) للعينات المترابطة.")
    t = between_groups(long_df, "بعدي", groups)
    if not t.empty:
        rep.table("يبين الفروق بين المجموعتين التجريبية والضابطة في الاختبار البعدي", t, note_sig + " اختبار (t) للعينات المستقلة.")
    t = improvement(long_df, groups)
    if not t.empty:
        rep.table("يبين نسب التطور لمجموعتي البحث", t, "نسبة التطور = (البعدي − القبلي) ÷ القبلي × 100، مع عكس الإشارة للمتغيرات التي يكون الأقل فيها أفضل.")
    return rep.bytes()


# ============================================================================
# 3) اتفاق الخبراء
# ============================================================================
RATINGS = ["يصلح", "لا يصلح", "يصلح بتعديل"]
FORMS = {"errors": "الأخطاء الفنية", "exercises": "التمرينات التصحيحية", "tests": "الاختبارات"}


def form_items(bank, form):
    if form == "errors":
        d = bank["errors"].sort_values("رمز_الخطأ")
        return pd.DataFrame({"الرمز": d["رمز_الخطأ"], "البند": d["الخطأ"]})
    if form == "exercises":
        d = bank["exercises"]
        return pd.DataFrame({"الرمز": d["رمز_التمرين"], "البند": d["التمرين"].astype(str)
                             .str.replace("{box}", "س").str.replace("{hurdle}", "س")})
    d = bank["tests"]
    return pd.DataFrame({"الرمز": d["رمز_الاختبار"], "البند": d["الاختبار"]})


def agreement_template(bank, form, n_experts):
    items = form_items(bank, form)
    for i in range(1, n_experts + 1):
        items[f"خبير {i}"] = ""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        items.to_excel(w, sheet_name="الإجابات", index=False)
        pd.DataFrame({"تعليمات": ["اكتب لكل خبير: يصلح أو لا يصلح أو يصلح بتعديل.",
                                  "لا تغيّر عمود الرمز.", "يمكن إضافة أعمدة خبراء أخرى بنفس الصيغة (خبير 6 ...)."]}
                     ).to_excel(w, sheet_name="اقرأني", index=False)
    wb_bytes = buf.getvalue()
    # قائمة منسدلة للتقييم
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    ws = wb["الإجابات"]
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["B"].width = 55
    dv = DataValidation(type="list", formula1='"يصلح,لا يصلح,يصلح بتعديل"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"C2:{openpyxl.utils.get_column_letter(2 + n_experts + 10)}{len(items) + 1}")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def agreement(answers, threshold=75.0, modify_counts=True):
    """answers: جدول فيه الرمز والبند وأعمدة الخبراء."""
    ex_cols = [c for c in answers.columns if c not in ("الرمز", "البند")]
    rows = []
    for _, r in answers.iterrows():
        vals = [str(r[c]).strip() for c in ex_cols if str(r[c]).strip() in RATINGS]
        n = len(vals)
        yes = sum(v == "يصلح" or (modify_counts and v == "يصلح بتعديل") for v in vals)
        mod = sum(v == "يصلح بتعديل" for v in vals)
        no = n - yes
        pct = yes / n * 100 if n else np.nan
        chi = p = np.nan
        if n:
            chi = ((yes - n / 2) ** 2 + (no - n / 2) ** 2) / (n / 2)
            p = 1 - stats.chi2.cdf(chi, 1)
        rows.append({"الرمز": r["الرمز"], "البند": r["البند"], "عدد الخبراء": n, "الموافقون": yes, "منهم بتعديل": mod,
                     "غير الموافقين": no, "النسبة المئوية": pct, "قيمة كا²": chi, "مستوى الدلالة": p,
                     "القرار": "—" if not n else ("مقبول" if pct >= threshold else "مرفوض")})
    return pd.DataFrame(rows)


def agreement_word(res, form, threshold):
    rep = WordReport(landscape=True)
    d = res.copy()
    d.insert(0, "ت", range(1, len(d) + 1))
    d["النسبة المئوية"] = d["النسبة المئوية"].map(lambda v: _f(v, 1) + "%" if pd.notna(v) else "—")
    d["قيمة كا²"] = d["قيمة كا²"].map(lambda v: _f(v, 2))
    d["مستوى الدلالة"] = d["مستوى الدلالة"].map(lambda v: _f(v, 3))
    rep.table(f"يبين نسب اتفاق الخبراء وقيم (كا²) على {FORMS[form]}",
              d[["ت", "البند", "عدد الخبراء", "الموافقون", "غير الموافقين", "النسبة المئوية", "قيمة كا²", "مستوى الدلالة", "القرار"]],
              f"يُقبل البند عند نسبة اتفاق ({threshold:.0f}%) فأكثر. درجة الحرية (1)، مستوى الدلالة (0.05).")
    return rep.bytes()


def apply_to_bank(bank_bytes, form, res):
    """يحذف البنود المرفوضة من ملف الأنموذج (مع تمارين الأخطاء المرفوضة) ويعيد الملف الجديد."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(bank_bytes))
    rejected = set(res[res["القرار"] == "مرفوض"]["الرمز"].astype(str))
    sheet, key = {"errors": ("الأخطاء", "رمز_الخطأ"), "exercises": ("التمارين", "رمز_التمرين"),
                  "tests": ("الاختبارات", "رمز_الاختبار")}[form]

    def drop(ws, col_name, codes):
        head = [c.value for c in ws[1]]
        ci = head.index(col_name) + 1
        n = 0
        for r in range(ws.max_row, 1, -1):
            if str(ws.cell(r, ci).value) in codes:
                ws.delete_rows(r)
                n += 1
        return n

    n = drop(wb[sheet], key, rejected)
    extra = 0
    if form == "errors" and rejected:
        extra = drop(wb["التمارين"], "رمز_الخطأ", rejected)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), n, extra
